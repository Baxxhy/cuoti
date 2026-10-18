#!/usr/bin/env python3
from __future__ import annotations

import json
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASE = "http://127.0.0.1:8000/api/v1"


def unpack(response: httpx.Response):
    response.raise_for_status(); body = response.json(); assert body["success"], body
    return body["data"]


def main() -> int:
    client = httpx.Client(timeout=120)
    user = unpack(client.post(BASE + "/auth/login", json={"email": "e2e@minddimension.local", "password": "123456"}))
    headers = {"Authorization": f"Bearer {user['access_token']}"}
    candidates = unpack(client.get(BASE + "/papers/candidates", headers=headers,
                                   params={"mastery_min": 0, "mastery_max": 1}))
    assert len(candidates) >= 3
    chosen = candidates[:3]
    paper = unpack(client.post(BASE + "/papers", headers=headers, json={
        "title": "最终验收专项练习", "question_ids": [item["id"] for item in chosen],
        "items": [{"question_id": item["id"], "score": 8 + index * 2} for index, item in enumerate(chosen)],
        "filters": {"mastery_min": 0, "mastery_max": 1},
    }))
    detail = unpack(client.get(BASE + f"/papers/{paper['id']}", headers=headers))
    assert len(detail["items"]) == 3 and [item["score"] for item in detail["items"]] == [8, 10, 12]
    reversed_items = [{"question_id": item["question_id"], "score": 15 + index}
                      for index, item in enumerate(reversed(detail["items"]))]
    updated = unpack(client.patch(BASE + f"/papers/{paper['id']}", headers=headers,
                                  json={"title": "最终验收专项练习（已排序）", "items": reversed_items}))
    assert [item["question_id"] for item in updated["items"]] == [item["question_id"] for item in reversed(detail["items"])]
    assert [item["score"] for item in updated["items"]] == [15, 16, 17]
    files: dict[str, dict[str, object]] = {}
    for kind in ("question", "answer", "excel"):
        job = unpack(client.post(BASE + f"/papers/{paper['id']}/export", headers=headers, json={"kind": kind}))
        path = ROOT / job["path"]
        assert path.exists() and path.stat().st_size > 500
        if kind == "excel":
            workbook = load_workbook(path, read_only=True); sheet = workbook["学习数据"]
            assert sheet.max_row == 4 and sheet.max_column == 5
        else:
            assert path.read_bytes()[:4] == b"%PDF"
        files[kind] = {"path": str(path), "bytes": path.stat().st_size}
    final_detail = unpack(client.get(BASE + f"/papers/{paper['id']}", headers=headers))
    assert len(final_detail["exports"]) >= 3
    disposable = unpack(client.post(BASE + "/papers", headers=headers, json={
        "title": "删除功能验收", "question_ids": [chosen[0]["id"]],
    }))
    assert unpack(client.delete(BASE + f"/papers/{disposable['id']}", headers=headers))["deleted"]
    report = {"passed": True, "candidate_count": len(candidates), "paper_id": paper["id"],
              "item_count": 3, "scores": [15, 16, 17], "reorder_verified": True,
              "delete_verified": True, "export_history_count": len(final_detail["exports"]), "files": files}
    target = ROOT / "reports/verification/papers.json"; target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False)); return 0


if __name__ == "__main__":
    raise SystemExit(main())
