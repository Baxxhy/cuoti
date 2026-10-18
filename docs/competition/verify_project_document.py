#!/usr/bin/env python3
"""Validate the competition manuscript and final Word deliverable."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document


HERE = Path(__file__).resolve().parent
MANIFEST = HERE / "figure_manifest.json"
OFFICIAL_HEADINGS = [
    "1 项目概况", "1.1 背景和基础", "1.2 场景和价值", "1.3 所需支持",
    "2 项目规划", "2.1 整体目标", "2.2 技术创新点",
    "3 实施方案", "3.1 技术可行性分析", "3.2 技术细节", "3.3 计划和分工",
    "4 参考资料",
]
FORBIDDEN = [
    "EduMind", "EduVision-Nano", "PaddleOCR", "Spring Boot", "微服务",
    "应用商店", "离线模式", "学校", "学院", "导师",
]
FRONT_MATTER_FORBIDDEN = ["摘要", "术语与阅读说明", "图表与证据使用约定"]
BOUNDARY_TERMS = ["当前原型结果", "目标指标", "后续预期"]


def load_manifest() -> list[dict]:
    data = json.loads(MANIFEST.read_text(encoding="utf-8"))
    assert len(data) == 24, f"图清单应有 24 项，实际 {len(data)} 项"
    ids = [item["id"] for item in data]
    expected = [*(f"1-{i}" for i in range(1, 7)), *(f"2-{i}" for i in range(1, 8)), *(f"3-{i}" for i in range(1, 12))]
    assert ids == expected, f"图号顺序不正确：{ids}"
    required = {"id", "chapter", "title", "kind", "width_cm", "height_cm", "output", "sources"}
    for item in data:
        missing = required - set(item)
        assert not missing, f"图 {item.get('id')} 缺少字段：{sorted(missing)}"
    return data


def normalize_heading(line: str) -> str:
    return re.sub(r"^#+\s*", "", line.strip()).replace("第一章 ", "1 ").replace("第二章 ", "2 ").replace("第三章 ", "3 ").replace("第四章 ", "4 ")


def source_checks(path: Path, manifest: list[dict]) -> list[str]:
    errors: list[str] = []
    text = path.read_text(encoding="utf-8")
    headings = [normalize_heading(line) for line in text.splitlines() if line.lstrip().startswith("#")]
    for heading in OFFICIAL_HEADINGS:
        if heading not in headings:
            errors.append(f"源稿缺少官方标题：{heading}")
    for heading in FRONT_MATTER_FORBIDDEN:
        if any(value == heading for value in headings):
            errors.append(f"源稿存在模板外前置章节：{heading}")
    if "占位" in text:
        errors.append("源稿仍包含图或结果占位文字")
    ids = re.findall(r"【图\s+(\d+-\d+)】", text)
    expected = [item["id"] for item in manifest]
    if ids != expected:
        errors.append(f"源稿图引用应严格为 24 个连续图号；实际 {ids}")
    for term in FORBIDDEN:
        if term in text:
            errors.append(f"源稿包含禁用或匿名风险词：{term}")
    for term in BOUNDARY_TERMS:
        if term not in text:
            errors.append(f"源稿缺少结果边界术语：{term}")
    if len(re.sub(r"\s+", "", text)) < 30_000:
        errors.append("源稿有效字符不足 30000，难以与 24 组图共同支撑完整竞赛文档")
    return errors


def docx_checks(path: Path, manifest: list[dict]) -> list[str]:
    errors: list[str] = []
    if not path.exists():
        return [f"Word 不存在：{path}"]
    try:
        document = Document(path)
    except Exception as exc:  # pragma: no cover - error reporting path
        return [f"Word 无法打开：{exc}"]
    text_parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    text = "\n".join(text_parts)
    for heading in OFFICIAL_HEADINGS:
        if heading not in text:
            errors.append(f"Word 缺少官方标题：{heading}")
    for term in FORBIDDEN:
        if term in text:
            errors.append(f"Word 包含禁用或匿名风险词：{term}")
    for term in BOUNDARY_TERMS:
        if term not in text:
            errors.append(f"Word 缺少结果边界术语：{term}")
    for item in manifest:
        caption = f"图 {item['id']}"
        if caption not in text:
            errors.append(f"Word 缺少题注：{caption}")
        marker = f"【图 {item['id']} 待替换】"
        if marker not in text:
            errors.append(f"Word 缺少规范插图说明框：{marker}")
        for source in item.get("sources", []):
            if source.endswith(".vue") and source not in text:
                errors.append(f"Word 图 {item['id']} 未标注 Vue 截图来源：{source}")
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        xml = archive.read("word/document.xml").decode("utf-8", "ignore")
        manual_breaks = len(re.findall(r'w:type="page"', xml))
        if not 57 <= manual_breaks <= 61:
            errors.append(f"Word 手工分页边界应支撑约 60 页，实际 {manual_breaks} 个")
        core = archive.read("docProps/core.xml").decode("utf-8", "ignore") if "docProps/core.xml" in names else ""
        for term in ["Larry", "学校", "学院", "导师"]:
            if term in core:
                errors.append(f"Word 元数据包含需清理信息：{term}")
    return errors


def pdf_checks(path: Path) -> list[str]:
    errors: list[str] = []
    if not path.exists():
        return [f"分页检查 PDF 不存在：{path}"]
    data = path.read_bytes()
    pages = len(re.findall(rb"/Type\s*/Page\b", data))
    if not 58 <= pages <= 62:
        errors.append(f"分页应为 58～62 页，实际 {pages} 页")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--docx", type=Path)
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--source-only", action="store_true")
    args = parser.parse_args()
    manifest = load_manifest()
    errors = source_checks(args.source, manifest)
    if not args.source_only:
        if not args.docx:
            errors.append("非 source-only 模式必须提供 --docx")
        else:
            errors.extend(docx_checks(args.docx, manifest))
    if args.pdf:
        errors.extend(pdf_checks(args.pdf))
    if errors:
        print("VALIDATION FAILED")
        for error in errors:
            print(f"- {error}")
        return 1
    print(f"VALIDATION PASSED: 24 figures, {len(OFFICIAL_HEADINGS)} official headings")
    return 0


if __name__ == "__main__":
    sys.exit(main())
